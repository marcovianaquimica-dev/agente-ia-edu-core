"""
Binary renderers (PDF via PyMuPDF, XLSX via openpyxl) for Teacher/Coordination
Portal report exports.

Background: ``report_export.ReportExportService`` already computes the REAL
report data (sourced from TeacherPortalService / CoordinationPortalService)
and shapes it into a structured payload dict - that part was correct. The bug
was purely in delivery: the routes returned that payload dict as JSON while
claiming ``content_type`` "application/pdf" or the XLSX MIME type. This
module turns the SAME payload into actual binary bytes for those two
formats. It invents no report data of its own - only presentation.

PDF -> PyMuPDF (``fitz``), same page-flow pattern already used by
``services.list_export.render_pdf`` (project precedent).
XLSX -> ``openpyxl`` (newly added dependency; no project precedent existed).
"""

from __future__ import annotations

import io
from typing import Any

_A4 = (595.0, 842.0)  # points, matches services.list_export
_MARGIN = 56.0
_LINE = 14.0

# Keys that describe the export itself, not report content - never rendered
# as a body section (title/generated_at ARE used, but as the header, not a
# section row).
_METADATA_KEYS = {"export_format", "filename", "content_type", "title", "generated_at"}

_LABELS: dict[str, str] = {
    "student_count": "Total de Alunos",
    "active_students_count": "Alunos Ativos",
    "overall_class_average": "Média Geral da Turma",
    "struggling_percentage": "% Com Dificuldade",
    "developing_percentage": "% Em Desenvolvimento",
    "mastered_percentage": "% Consolidado",
    "accuracy_percentage": "Taxa de Acerto (%)",
    "total_questions_answered": "Questões Respondidas",
    "total_questions_correct": "Questões Corretas",
    "content_node_id": "ID do Conteúdo",
    "content_name": "Conteúdo",
    "class_average_mastery": "Média da Turma (%)",
    "students_struggling_count": "Alunos com Dificuldade",
    "struggling_students_count": "Alunos com Dificuldade",
    "total_students": "Total de Alunos",
    "reason": "Motivo",
    "priority": "Prioridade",
    "impacted_students_count": "Alunos Impactados",
    "evidence": "Evidência",
    "recommended_action": "Ação Recomendada",
    "last_lesson_date": "Última Aula",
    "teacher_or_author": "Professor/Autor",
    "student_id": "Aluno",
    "name": "Nome",
    "classroom_id": "Turma",
    "average_mastery": "Domínio Médio (%)",
    "status_label": "Status",
    "mastery_score": "Domínio (%)",
    "current_level": "Nível",
    "questions_answered": "Questões Respondidas",
    "questions_correct": "Questões Corretas",
}

_SECTION_TITLES: dict[str, str] = {
    "mastery_distribution": "Distribuição de Domínio",
    "strengths": "Pontos Fortes",
    "improvement_areas": "Áreas de Melhoria",
    "recent_contents_taught": "Conteúdos Recentes",
    "action_plan": "Plano de Ação",
    "students_roster": "Alunos",
    "content_masteries": "Domínio por Conteúdo",
    "priority_contents": "Conteúdos Prioritários",
    "current_recommendations": "Recomendações Atuais",
}

# Order in which list-shaped payload keys become sections, when present.
_LIST_SECTION_ORDER = [
    "strengths",
    "improvement_areas",
    "recent_contents_taught",
    "action_plan",
    "students_roster",
    "content_masteries",
    "priority_contents",
    "current_recommendations",
]


def _label(key: str) -> str:
    return _LABELS.get(key, key.replace("_", " ").strip().capitalize())


def _fmt_scalar(key: str, value: Any) -> str:
    if value is None:
        return "-"
    if isinstance(value, bool):
        return "Sim" if value else "Não"
    if isinstance(value, float):
        return f"{value:.1f}"
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    return str(value)


def build_sections(payload: dict[str, Any]) -> list[dict[str, Any]]:
    """Turns a ReportExportService payload dict into an ordered list of
    render-agnostic sections. Both render_pdf and render_xlsx consume this
    SAME structure, so the two formats can never drift apart from one
    another or from what export_classroom_report/export_student_report
    actually computed.
    """
    sections: list[dict[str, Any]] = []

    resumo: dict[str, Any] = {}
    summary_val = payload.get("summary")
    if isinstance(summary_val, dict):
        resumo.update(summary_val)
    for scalar_key in ("accuracy_percentage", "total_questions_answered", "total_questions_correct"):
        if scalar_key in payload and payload[scalar_key] is not None:
            resumo.setdefault(scalar_key, payload[scalar_key])
    if resumo:
        sections.append({"title": "Resumo", "kind": "kv", "data": resumo})

    mdist = payload.get("mastery_distribution")
    if isinstance(mdist, dict) and mdist:
        sections.append({"title": _SECTION_TITLES["mastery_distribution"], "kind": "kv", "data": mdist})

    for key in _LIST_SECTION_ORDER:
        items = payload.get(key)
        if not items:
            continue
        if isinstance(items[0], dict):
            columns: list[str] = []
            for it in items:
                for k in it.keys():
                    if k not in columns:
                        columns.append(k)
            rows = [[_fmt_scalar(c, it.get(c)) for c in columns] for it in items]
            sections.append({
                "title": _SECTION_TITLES.get(key, key),
                "kind": "table",
                "columns": [_label(c) for c in columns],
                "rows": rows,
            })
        else:
            sections.append({
                "title": _SECTION_TITLES.get(key, key),
                "kind": "list",
                "data": [str(x) for x in items],
            })
    return sections


# ---------------------------------------------------------------------------
# PDF (PyMuPDF)
# ---------------------------------------------------------------------------


def pdf_available() -> bool:
    try:
        import pymupdf  # noqa: F401
        return True
    except Exception:
        return False


def render_pdf(payload: dict[str, Any]) -> bytes:
    import pymupdf as fitz

    doc = fitz.open()
    state: dict[str, Any] = {"page": None, "y": 0.0}

    def new_page():
        state["page"] = doc.new_page(width=_A4[0], height=_A4[1])
        state["y"] = _MARGIN

    def _wrap(text: str, max_w: float, size: float, font: str) -> list[str]:
        out: list[str] = []
        for para in (text or "").split("\n"):
            words = para.split(" ")
            cur = ""
            for w in words:
                trial = (cur + " " + w).strip()
                if fitz.get_text_length(trial, fontname=font, fontsize=size) <= max_w or not cur:
                    cur = trial
                else:
                    out.append(cur)
                    cur = w
            out.append(cur)
        return out or [""]

    def write(text: str, *, size: float = 10.0, bold: bool = False, gap: float = 0.0,
              color=(0, 0, 0), indent: float = 0.0):
        if state["page"] is None:
            new_page()
        state["y"] += gap
        font = "helv" if not bold else "hebo"
        max_w = _A4[0] - 2 * _MARGIN - indent
        for line in _wrap(text, max_w, size, font):
            if state["y"] > _A4[1] - _MARGIN:
                new_page()
            state["page"].insert_text((_MARGIN + indent, state["y"]), line,
                                       fontsize=size, fontname=font, color=color)
            state["y"] += _LINE

    new_page()
    write(payload.get("title") or "Relatório", size=16, bold=True, gap=4)
    generated_at = payload.get("generated_at") or ""
    write(f"Gerado em {generated_at}", size=8, color=(0.4, 0.4, 0.4), gap=2)
    write("", gap=6)

    for section in build_sections(payload):
        write(section["title"], size=13, bold=True, gap=16)
        if section["kind"] == "kv":
            for k, v in section["data"].items():
                write(f"{_label(k)}: {_fmt_scalar(k, v)}", size=10, gap=2)
        elif section["kind"] == "list":
            for item in section["data"]:
                write(f"- {item}", size=10, indent=10, gap=2)
        elif section["kind"] == "table":
            cols = section["columns"]
            for i, row in enumerate(section["rows"], start=1):
                line = "; ".join(f"{c}: {v}" for c, v in zip(cols, row))
                write(f"{i}. {line}", size=9, indent=10, gap=3)

    for i, pg in enumerate(doc, start=1):
        pg.insert_text((_A4[0] - _MARGIN - 70, _A4[1] - 30),
                        f"Página {i} de {len(doc)}", fontsize=7, color=(0.4, 0.4, 0.4))

    out = doc.tobytes()
    doc.close()
    return out


# ---------------------------------------------------------------------------
# XLSX (openpyxl)
# ---------------------------------------------------------------------------

_INVALID_SHEET_CHARS = set("[]:*?/\\")


def _sheet_title(name: str, used: set[str]) -> str:
    cleaned = "".join(c for c in name if c not in _INVALID_SHEET_CHARS).strip() or "Sheet"
    cleaned = cleaned[:31]
    base = cleaned
    n = 2
    while cleaned in used:
        suffix = f" ({n})"
        cleaned = (base[: 31 - len(suffix)] + suffix)
        n += 1
    used.add(cleaned)
    return cleaned


def render_xlsx(payload: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    used_titles: set[str] = set()

    cover = wb.active
    cover.title = _sheet_title("Relatorio", used_titles)
    cover["A1"] = payload.get("title") or "Relatório"
    cover["A1"].font = Font(bold=True, size=14)
    cover["A2"] = f"Gerado em {payload.get('generated_at') or ''}"

    for section in build_sections(payload):
        ws = wb.create_sheet(_sheet_title(section["title"], used_titles))
        ws["A1"] = section["title"]
        ws["A1"].font = Font(bold=True, size=12)

        if section["kind"] == "kv":
            r = 3
            for k, v in section["data"].items():
                ws.cell(row=r, column=1, value=_label(k))
                ws.cell(row=r, column=2, value=v if isinstance(v, (int, float)) else _fmt_scalar(k, v))
                r += 1
        elif section["kind"] == "list":
            r = 3
            for item in section["data"]:
                ws.cell(row=r, column=1, value=item)
                r += 1
        elif section["kind"] == "table":
            for c_idx, col_name in enumerate(section["columns"], start=1):
                cell = ws.cell(row=3, column=c_idx, value=col_name)
                cell.font = Font(bold=True)
            for r_idx, data_row in enumerate(section["rows"], start=4):
                for c_idx, value in enumerate(data_row, start=1):
                    ws.cell(row=r_idx, column=c_idx, value=value)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = ["build_sections", "pdf_available", "render_pdf", "render_xlsx"]
