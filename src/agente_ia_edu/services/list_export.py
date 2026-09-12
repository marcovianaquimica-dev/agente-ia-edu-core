"""PHASE 15 - PDF and DOCX export of a persisted question list.

Both renderers consume the SAME :class:`GeneratedListDefinition` (as a dict) via
one shared render model - there is a single prova-assembly path, not two.

Rules:
  * Official statement + options A-E verbatim, in the persisted order.
  * Answer key rendered only when the stored presentation includes it.
  * Resolution rendered only when it is actually available. It never is in the
    current schema, so mode 3 prints "Resolução não disponível." - no invented text.

PDF -> PyMuPDF (``fitz``), already adopted by the project (text-recovery extra).
DOCX -> ``python-docx``. If PyMuPDF is not installed the caller should surface a
clean 503 rather than crash (see :func:`pdf_available`).
"""

from __future__ import annotations

import io
from typing import Any

_A4 = (595.0, 842.0)  # points
_MARGIN = 56.0
_LINE = 13.5
_RESOLUTION_UNAVAILABLE = "Resolução não disponível."


def pdf_available() -> bool:
    try:
        import pymupdf  # noqa: F401
        return True
    except Exception:
        return False


# ---------------------------------------------------------------------------
# shared render model
# ---------------------------------------------------------------------------


def build_render_model(definition: dict) -> dict:
    cfg = definition["configuration"]
    include_key = bool(definition.get("answer_key_included"))
    include_resolution = bool(definition.get("resolution_included"))
    items = []
    key_rows = []
    for it in definition["items"]:
        items.append({
            "position": it["position"],
            "heading": _question_heading(it),
            "statement": (it.get("statement") or "").strip(),
            "options": [(o["key"], (o["text"] or "").strip())
                        for o in sorted(it["options"], key=lambda x: x["position"])],
        })
        if include_key:
            ak = it.get("answer_key") or {}
            res = ak.get("resolution") or {}
            resolution_text = None
            if include_resolution:
                resolution_text = res.get("text") if res.get("available") else _RESOLUTION_UNAVAILABLE
            key_rows.append({
                "position": it["position"],
                "correct_option_key": ak.get("correct_option_key", "?"),
                "resolution_text": resolution_text,
            })
    return {
        "title": cfg["title"],
        "instructions": (cfg.get("instructions") or "").strip() or None,
        "activity_mode": cfg.get("activity_mode", "EXERCISE_LIST"),
        "question_count": definition["question_count"],
        "answer_key_presentation": cfg.get("answer_key_presentation", "NONE"),
        "resolution_style": cfg.get("resolution_style"),
        "include_key": include_key,
        "include_resolution": include_resolution,
        "items": items,
        "key_rows": key_rows,
        "fingerprint": definition.get("selection_fingerprint", ""),
        "generated_at": definition.get("generated_at", ""),
    }


def _question_heading(it: dict) -> str:
    bits = [f"Questão {it['position']}"]
    src = []
    if it.get("source"):
        src.append(str(it["source"]))
    if it.get("year"):
        src.append(str(it["year"]))
    if it.get("official_number"):
        src.append(f"Q{it['official_number']}")
    if it.get("enem_area"):
        src.append(str(it["enem_area"]))
    if src:
        bits.append("(" + " · ".join(src) + ")")
    return " ".join(bits)


# ---------------------------------------------------------------------------
# PDF (PyMuPDF)
# ---------------------------------------------------------------------------


def render_pdf(model: dict) -> bytes:
    import pymupdf as fitz

    doc = fitz.open()
    state = {"page": None, "y": 0.0, "n": 0}

    def new_page():
        state["n"] += 1
        state["page"] = doc.new_page(width=_A4[0], height=_A4[1])
        state["y"] = _MARGIN
        state["page"].insert_text(
            (_MARGIN, _A4[1] - _MARGIN + 20), f"AGENTE IA EDU · {model['title'][:70]}",
            fontsize=7, color=(0.4, 0.4, 0.4),
        )

    def _footer_page_numbers():
        for i, pg in enumerate(doc, start=1):
            pg.insert_text((_A4[0] - _MARGIN - 60, _A4[1] - _MARGIN + 20),
                           f"Página {i} de {len(doc)}", fontsize=7, color=(0.4, 0.4, 0.4))

    def write(text: str, *, size: float = 10.0, bold: bool = False, gap: float = 0.0,
              color=(0, 0, 0), indent: float = 0.0):
        if state["page"] is None:
            new_page()
        state["y"] += gap
        font = "helv" if not bold else "hebo"
        max_w = _A4[0] - 2 * _MARGIN - indent
        for line in _wrap(text, max_w, size, font):
            if state["y"] > _A4[1] - _MARGIN - _LINE:
                new_page()
            state["page"].insert_text((_MARGIN + indent, state["y"]), line,
                                      fontsize=size, fontname=font, color=color)
            state["y"] += _LINE

    def _wrap(text: str, max_w: float, size: float, font: str) -> list[str]:
        import pymupdf as _f
        out: list[str] = []
        for para in (text or "").split("\n"):
            words = para.split(" ")
            cur = ""
            for w in words:
                trial = (cur + " " + w).strip()
                if _f.get_text_length(trial, fontname=font, fontsize=size) <= max_w or not cur:
                    cur = trial
                else:
                    out.append(cur)
                    cur = w
            out.append(cur)
        return out or [""]

    new_page()
    write(model["title"], size=16, bold=True, gap=4)
    write(f"{model['question_count']} questões · {model['answer_key_presentation']}"
          + (f" · {model['resolution_style']}" if model["include_resolution"] else ""),
          size=8, color=(0.35, 0.35, 0.35), gap=2)
    if model["instructions"]:
        write("Instruções: " + model["instructions"], size=9, gap=8, color=(0.15, 0.15, 0.15))
    write("", gap=6)

    for q in model["items"]:
        write(q["heading"], size=11, bold=True, gap=12)
        write(q["statement"], size=10, gap=3)
        for key, txt in q["options"]:
            write(f"{key}) {txt}", size=10, indent=14, gap=1)

    if model["include_key"]:
        write("GABARITO", size=13, bold=True, gap=20)
        for row in model["key_rows"]:
            write(f"{row['position']}. {row['correct_option_key']}", size=10, gap=1)
            if model["include_resolution"]:
                write("Resolução: " + (row["resolution_text"] or _RESOLUTION_UNAVAILABLE),
                      size=9, indent=14, color=(0.3, 0.3, 0.3), gap=1)

    _footer_page_numbers()
    out = doc.tobytes()
    doc.close()
    return out


# ---------------------------------------------------------------------------
# DOCX (python-docx)
# ---------------------------------------------------------------------------


def render_docx(model: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.core_properties.title = model["title"]

    h = doc.add_heading(model["title"], level=0)
    meta = doc.add_paragraph(
        f"{model['question_count']} questões · {model['answer_key_presentation']}"
        + (f" · {model['resolution_style']}" if model["include_resolution"] else "")
    )
    meta.runs[0].font.size = Pt(8)
    if model["instructions"]:
        p = doc.add_paragraph()
        r = p.add_run("Instruções: ")
        r.bold = True
        p.add_run(model["instructions"])

    for q in model["items"]:
        qh = doc.add_paragraph()
        qh.add_run(q["heading"]).bold = True
        doc.add_paragraph(q["statement"])
        for key, txt in q["options"]:
            op = doc.add_paragraph(style="List Bullet")
            op.add_run(f"{key}) {txt}")

    if model["include_key"]:
        doc.add_page_break()
        doc.add_heading("Gabarito", level=1)
        for row in model["key_rows"]:
            kp = doc.add_paragraph()
            kp.add_run(f"{row['position']}. ").bold = True
            kp.add_run(row["correct_option_key"])
            if model["include_resolution"]:
                rp = doc.add_paragraph()
                rp.add_run("Resolução: ").italic = True
                rp.add_run(row["resolution_text"] or _RESOLUTION_UNAVAILABLE).italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_both(definition: dict) -> dict[str, Any]:
    """Convenience for tests: PDF + DOCX from ONE definition, same render model."""
    model = build_render_model(definition)
    result: dict[str, Any] = {"model": model, "docx": render_docx(model)}
    if pdf_available():
        result["pdf"] = render_pdf(model)
    return result


__all__ = ["build_render_model", "pdf_available", "render_both", "render_docx", "render_pdf"]
