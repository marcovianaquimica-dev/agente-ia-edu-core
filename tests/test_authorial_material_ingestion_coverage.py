"""Service-level coverage for src/agente_ia_edu/services/authorial_material_ingestion.py.

Wave-7 (concurrent-zone campaign). Targets branches left uncovered by
tests/test_authorial_ingestion_http.py (HTTP-route-layer, which duplicates the
route's OWN pre-checks and never reaches several of the service's own error
paths) and tests/test_phase26_authorial_material_ingestion.py (service-level
happy paths). Measured baseline just before this file (existing test files
only): 200 statements, 15 missing - 103, 127-128, 132, 134, 180-181, 273, 275,
281, 283, 298, 340, 390, 440.

Uses a real async SQLAlchemy session against SQLite with the session_factory
default (``expire_on_commit=True``, matching production's
``create_session_factory()`` in ``db/session.py``) - not the ``False`` override
several older test files use - since this is exactly the setting under which
the MissingGreenlet-after-commit bug class this campaign hunts for manifests.
"""

from __future__ import annotations

import base64
import unittest
import uuid as _uuid
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

from agente_ia_edu.db.base import Base
from agente_ia_edu.db.models import CatalogNode, MaterialBlock
from agente_ia_edu.services.authorial_material_ingestion import (
    AuthorialIngestionError,
    AuthorialMaterialIngestionService,
    _reload_content_lines,
)
from agente_ia_edu.services.material_storage import MaterialStorage

_TMP = Path("/tmp/authorial_ingestion_coverage_fixtures")
_STORAGE_ROOT = Path("/tmp/authorial_ingestion_coverage_storage")

# a minimal, valid 1x1 transparent PNG - just enough for python-docx's own
# header-sniffing image code (no PIL dependency in this environment) to
# accept it as a real embeddable picture.
_PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _make_docx(path: Path, *, title: str = "Apostila Teste", with_image: bool = False) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph(f"Episódio 01 – {title}")
    doc.add_paragraph(f"Um parágrafo de introdução sobre o tema estudado ({path.stem}).")
    doc.add_paragraph("Questão 1. Qual é a resposta correta?")
    doc.add_paragraph("a) Alternativa um")
    doc.add_paragraph("b) Alternativa dois")
    if with_image:
        png_path = path.with_suffix(".png")
        png_path.write_bytes(_PNG_1PX)
        doc.add_picture(str(png_path))
    doc.save(path)


class AuthorialMaterialIngestionCoverageTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self) -> None:
        _TMP.mkdir(exist_ok=True)
        self.engine = create_async_engine("sqlite+aiosqlite:///:memory:", poolclass=StaticPool)
        async with self.engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        # deliberately no expire_on_commit override - async_sessionmaker's own
        # default is True, matching create_session_factory() in production.
        self.factory = async_sessionmaker(self.engine, class_=AsyncSession)
        self.storage = MaterialStorage(root=_STORAGE_ROOT / _uuid.uuid4().hex)

        async with self.factory() as s:
            d = CatalogNode(code="AMIC-DISC", name="Disciplina Teste", node_type="DISCIPLINE", active=True)
            s.add(d)
            await s.flush()
            d.root_id = d.id
            a = CatalogNode(
                code="AMIC-AREA", name="Area Teste", node_type="AREA",
                parent_id=d.id, root_id=d.id, active=True,
            )
            s.add(a)
            await s.flush()
            c = CatalogNode(
                code="AMIC-CONTENT", name="Apostila Teste", node_type="CONTENT",
                parent_id=a.id, root_id=d.id, active=True,
            )
            s.add(c)
            await s.commit()

    async def asyncTearDown(self) -> None:
        await self.engine.dispose()

    def _svc(self, session: AsyncSession) -> AuthorialMaterialIngestionService:
        return AuthorialMaterialIngestionService(session, storage=self.storage)

    # -- line 103: ingest_file on a path that does not exist ----------------
    async def test_ingest_nonexistent_file_raises_extraction_failed(self):
        async with self.factory() as session:
            svc = self._svc(session)
            with self.assertRaises(AuthorialIngestionError) as ctx:
                await svc.ingest_file(
                    Path("/tmp/authorial_ingestion_coverage_fixtures/does_not_exist_xyz.pdf"),
                    uploaded_by="prof_a", school_id=None,
                )
            self.assertEqual(ctx.exception.code, "EXTRACTION_FAILED")
            self.assertIn("does not exist", str(ctx.exception))

    # -- lines 127-128: parse_authorial_document raises -> EXTRACTION_FAILED
    async def test_corrupt_docx_raises_extraction_failed(self):
        path = _TMP / "corrupt.docx"
        path.write_bytes(b"this is not a real docx/zip payload at all")
        async with self.factory() as session:
            svc = self._svc(session)
            with self.assertRaises(AuthorialIngestionError) as ctx:
                await svc.ingest_file(path, uploaded_by="prof_a", school_id=None)
            self.assertEqual(ctx.exception.code, "EXTRACTION_FAILED")

    # -- line 134: embedded images -> visual structure_issue, never dropped -
    async def test_docx_with_embedded_image_flags_visual_structure_issue(self):
        path = _TMP / "visual.docx"
        _make_docx(path, title="Apostila Teste", with_image=True)
        async with self.factory() as session:
            svc = self._svc(session)
            review, created = await svc.ingest_file(path, uploaded_by="prof_a", school_id=None)
            self.assertTrue(created)
            self.assertTrue(
                any("visual element" in issue for issue in (review.structure_issues or [])),
                review.structure_issues,
            )

    # -- lines 180-181: curriculum matcher raises -> CURRICULUM_MAPPING_FAILED
    async def test_curriculum_matcher_failure_is_reported_not_masked(self):
        path = _TMP / "matcher_boom.docx"
        _make_docx(path, title="Apostila Teste")
        async with self.factory() as session:
            svc = self._svc(session)

            async def _boom(_texts):
                raise RuntimeError("matcher backend unavailable")

            svc._matcher.match = _boom
            with self.assertRaises(AuthorialIngestionError) as ctx:
                await svc.ingest_file(path, uploaded_by="prof_a", school_id=None)
            self.assertEqual(ctx.exception.code, "CURRICULUM_MAPPING_FAILED")
            self.assertIn("matcher backend unavailable", str(ctx.exception))

    # -- lines 132 + 340: no section/chapter structure -> flagged, and a
    #    review that reaches APPROVED with zero sections is still blocked at
    #    publish time (spec s15: never publish a structurally-broken material)
    async def test_no_structure_document_is_flagged_and_blocks_publish_once_approved(self):
        path = _TMP / "no_structure.txt"
        # a single line with no blank-line-separated content after it: the
        # whole line is consumed as the document TITLE (matches the seeded
        # CONTENT node's name so classification auto-resolves to MAPPED) and
        # parse_authorial_text() never creates a single ParsedSection.
        path.write_text("Apostila Teste", encoding="utf-8")
        async with self.factory() as session:
            svc = self._svc(session)
            review, created = await svc.ingest_file(path, uploaded_by="prof_a", school_id=None)
            self.assertTrue(created)
            self.assertEqual(review.classification_state, "MAPPED")
            self.assertIn(
                "no section/chapter structure was detected", review.structure_issues or []
            )
            self.assertEqual(review.review_status, "NEEDS_REVIEW")

            approved = await svc.approve(review.id, reviewed_by="prof_a")
            self.assertEqual(approved.review_status, "APPROVED")

            with self.assertRaises(AuthorialIngestionError) as ctx:
                await svc.publish(review.id, published_by="prof_a")
            self.assertEqual(ctx.exception.code, "PUBLICATION_BLOCKED")
            self.assertIn("no structure to publish", str(ctx.exception))

    # -- lines 273, 275, 281, 283: update_classification's discipline_code /
    #    area_code / subcontent_codes / notes setters (only content_code was
    #    ever exercised by the existing HTTP/service test suites)
    async def test_update_classification_sets_discipline_area_subcontent_and_notes(self):
        path = _TMP / "classify_fields.docx"
        _make_docx(path, title="Zzqxw Unrelated Whatever 999")
        async with self.factory() as session:
            svc = self._svc(session)
            review, _ = await svc.ingest_file(path, uploaded_by="prof_a", school_id=None)
            self.assertEqual(review.classification_state, "TAXONOMY_GAP")

            updated = await svc.update_classification(
                review.id,
                discipline_code="AMIC-DISC",
                area_code="AMIC-AREA",
                subcontent_codes=["SUB-1", "SUB-2"],
                notes="revisão manual de metadados",
                reviewed_by="prof_a",
            )
            self.assertEqual(updated.discipline_code, "AMIC-DISC")
            self.assertEqual(updated.area_code, "AMIC-AREA")
            self.assertEqual(updated.subcontent_codes, ["SUB-1", "SUB-2"])
            self.assertEqual(updated.notes, "revisão manual de metadados")

    # -- line 298: approve() from a status other than PENDING_REVIEW/NEEDS_REVIEW
    async def test_approve_twice_is_blocked_on_the_second_call(self):
        path = _TMP / "approve_twice.docx"
        _make_docx(path, title="Apostila Teste")
        async with self.factory() as session:
            svc = self._svc(session)
            review, _ = await svc.ingest_file(path, uploaded_by="prof_a", school_id=None)
            await svc.approve(review.id, reviewed_by="prof_a")

            with self.assertRaises(AuthorialIngestionError) as ctx:
                await svc.approve(review.id, reviewed_by="prof_a")
            self.assertEqual(ctx.exception.code, "PUBLICATION_BLOCKED")
            self.assertIn("cannot approve from status", str(ctx.exception))

    # -- line 390: a "## " content line becomes a HEADING block on publish --
    async def test_publish_turns_numbered_subsection_into_heading_block(self):
        import fitz

        path = _TMP / "subsection.pdf"
        d = fitz.open()
        page = d.new_page()
        page.insert_text((72, 72), "TEMPORADA 1: TESTE DE INGESTAO", fontsize=14)
        # the trailing "." on its own text line closes the episode heading
        # right there (via _HEADING_LINE_END) so the "1.1 ..." numbered
        # subsection marker below lands in the section BODY, not swallowed
        # into the heading title itself.
        page.insert_text((72, 100), "EPISÓDIO 01 – Soluções Químicas.", fontsize=12)
        page.insert_text(
            (72, 130),
            "1.1 Conceito Fundamental. Uma solução química é uma mistura homogênea de substâncias.",
            fontsize=10,
        )
        d.save(str(path))

        async with self.factory() as session:
            svc = self._svc(session)
            review, created = await svc.ingest_file(path, uploaded_by="prof_a", school_id=None)
            self.assertTrue(created)

            await svc.update_classification(review.id, content_code="AMIC-CONTENT", reviewed_by="prof_a")
            await svc.approve(review.id, reviewed_by="prof_a")
            published = await svc.publish(review.id, published_by="prof_a")
            self.assertEqual(published.review_status, "PUBLISHED")

            headings = (await session.execute(
                select(MaterialBlock).where(MaterialBlock.block_type == "HEADING")
            )).scalars().all()
            self.assertTrue(
                any("Conceito Fundamental" in (h.title or "") for h in headings),
                [h.title for h in headings],
            )

    # -- line 440: _reload_content_lines' documented defensive fallback for a
    #    row that "somehow has no content_text" (a row created before PHASE
    #    26's additive content_text column-populate step existed, or by a
    #    different pathway entirely) - unreachable through the live
    #    ingest_file()->publish() flow today (content_text and
    #    content_preview are always populated together, from the same
    #    parsed_section.content_lines, in the same request), so exercised
    #    directly against the module-level helper per the function's own
    #    contract instead of inventing an artificial end-to-end path for it.
    def test_reload_content_lines_falls_back_to_preview_when_text_missing(self):
        from agente_ia_edu.db.models import IngestionSection

        legacy_row = IngestionSection(content_text=None, content_preview="resumo de uma seção antiga")
        self.assertEqual(_reload_content_lines(legacy_row), ["resumo de uma seção antiga"])

        empty_row = IngestionSection(content_text=None, content_preview=None)
        self.assertEqual(_reload_content_lines(empty_row), [])


if __name__ == "__main__":
    unittest.main()
