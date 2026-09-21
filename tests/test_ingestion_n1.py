"""Ingestion perf audit - N+1 regression test for
IngestionService.ingest_document().

ingest_document() used to persist sections and questions with a
`session.add(row); await session.flush()` PER ROW inside two separate loops,
issuing one INSERT round-trip per section AND per question instead of one
batched INSERT per collection (the discipline
question_extraction_service.run_extraction already documents and follows -
"Batched throughout ... never one query per question"). Confirmed LIVE
against real Postgres (port 5433, a synthetic .docx built with python-docx,
temporary IngestionDocument, cleaned up afterwards): a 5-question document
issued 12 SQL statements, a 50-question one issued 61 - real O(n) growth.
After the fix (client-side uuid4 ids assigned at construction so section/
question ids are known before any INSERT runs, then one `add_all()` + one
`flush()` per collection) both cases issue 7.

This test pins that down with a fast, in-process, disposable sqlite counter:
the query count must stay flat as the number of questions grows.
"""

from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from sqlalchemy import event
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

from agente_ia_edu.db.base import Base
from agente_ia_edu.services.ingestion import IngestionService


class QueryCounter:
    """Same `before_cursor_execute` technique as
    tests/test_phase24_study_session_n1.py / test_phase29_publish_run_n1.py."""

    def __init__(self, engine):
        self.engine = engine
        self.count = 0
        self._listener = None

    def __enter__(self):
        def _count(conn, cursor, statement, parameters, context, executemany):
            self.count += 1
        self._listener = _count
        event.listen(self.engine.sync_engine, "before_cursor_execute", self._listener)
        return self

    def __exit__(self, *exc):
        event.remove(self.engine.sync_engine, "before_cursor_execute", self._listener)


def _build_docx(path: Path, n_questions: int, *, tag: str) -> None:
    doc = Document()
    doc.add_paragraph(f"Temporada 1 - N1 {tag}")
    per_episode = 10
    episode = 0
    for i in range(1, n_questions + 1):
        if (i - 1) % per_episode == 0:
            episode += 1
            doc.add_paragraph(f"Episódio {episode} - N1 {tag}")
        doc.add_paragraph(f"Questão {i}. Enunciado da questao {i} {tag}?")
        for letter in "abcd":
            doc.add_paragraph(f"{letter}) alternativa {letter} da questao {i}")
    doc.save(str(path))


class IngestDocumentN1Tests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.engine = create_async_engine(
            "sqlite+aiosqlite:///:memory:", echo=False, poolclass=StaticPool)
        async with self.engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        self.factory = async_sessionmaker(self.engine, class_=AsyncSession, expire_on_commit=False)
        self._tmp = TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    async def asyncTearDown(self):
        await self.engine.dispose()
        self._tmp.cleanup()

    async def _ingest_and_count(self, n_questions: int, tag: str) -> tuple[int, int]:
        path = self.tmp / f"n1-{tag}.docx"
        _build_docx(path, n_questions, tag=tag)
        async with self.factory() as session:
            service = IngestionService()
            with QueryCounter(self.engine) as counter:
                document, run = await service.ingest_document(session, path, ingested_by="tester")
            self.assertEqual(run.questions_found, n_questions, run.questions_found)
            return counter.count, run.questions_found

    async def test_ingest_document_query_count_does_not_scale_with_question_count(self):
        count_small, _ = await self._ingest_and_count(5, "small")
        count_big, _ = await self._ingest_and_count(50, "big")

        # Before the fix this grew ~1 query per extra section/question (45
        # extra questions -> +49 SQL statements: 12 -> 61 measured live).
        # After the fix both cases issue the exact same small, fixed count.
        self.assertEqual(
            count_small, count_big,
            f"query count grew with question count ({count_small} -> {count_big}) - "
            "ingest_document's section/question persistence regressed to one flush per row",
        )


if __name__ == "__main__":
    unittest.main()
