"""HTTP-layer tests for teacher_portal_router (src/agente_ia_edu/api/routes/teacher_portal.py).

These exercise the FastAPI route handlers themselves - request parsing, the
ScopeAuthorizationError -> HTTPException(403) mapping, and the ValueError ->
HTTPException(400) mapping on export - which the existing test suite never
touched: every prior test in this domain (tests/test_teacher_portal.py,
tests/test_r0_*teacher*.py) calls TeacherPortalService directly, skipping the
HTTP layer entirely.
"""

import asyncio
import unittest
from uuid import uuid4

from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

from agente_ia_edu.api.dependencies import get_current_identity, get_session_factory
from agente_ia_edu.api.routes.teacher_portal import teacher_portal_router
from agente_ia_edu.db.base import Base
from agente_ia_edu.db.models import CatalogNode, StudentContentMastery
from agente_ia_edu.identity import ExternalIdentityContext
from agente_ia_edu.services.admin import AdminRole, AdminScopeType, PlatformAdminService


class TeacherPortalHTTP(unittest.TestCase):
    def setUp(self):
        async def setup():
            engine = create_async_engine("sqlite+aiosqlite:///:memory:", poolclass=StaticPool)
            async with engine.begin() as connection:
                await connection.run_sync(Base.metadata.create_all)
            factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
            async with factory() as session:
                admin_service = PlatformAdminService(session)
                school_a = await admin_service.create_school(performed_by_external_id="admin:master", code="TPH-A", name="Escola A")
                school_b = await admin_service.create_school(performed_by_external_id="admin:master", code="TPH-B", name="Escola B")

                content = CatalogNode(node_type="DISCIPLINE", name="Quimica", position=1, active=True)
                session.add(content)
                await session.flush()
                content.root_id = content.id
                topic = CatalogNode(parent_id=content.id, root_id=content.id, node_type="CONTENT", name="Diluicao", position=1, active=True)
                session.add(topic)
                await session.flush()

                await admin_service.link_user_to_school(
                    performed_by_external_id="admin:master", external_user_id="teacher-a",
                    role=AdminRole.TEACHER, scope_type=AdminScopeType.CLASSROOM,
                    school_id=school_a.id, scope_external_id="TURMA_3A",
                )
                await admin_service.link_user_to_school(
                    performed_by_external_id="admin:master", external_user_id="student-a",
                    role=AdminRole.STUDENT, scope_type=AdminScopeType.CLASSROOM,
                    school_id=school_a.id, scope_external_id="TURMA_3A",
                )
                session.add(StudentContentMastery(external_identity_id="student-a", content_node_id=topic.id, mastery_score=40.0))
                await session.commit()
                return engine, factory, school_a.id, school_b.id

        self.engine, self.session_factory, self.school_a, self.school_b = asyncio.run(setup())
        self.identity = {"value": ExternalIdentityContext(provider="test", external_user_id="teacher-a", roles=("teacher",))}
        app = FastAPI()
        app.include_router(teacher_portal_router)
        app.dependency_overrides[get_session_factory] = lambda: self.session_factory
        app.dependency_overrides[get_current_identity] = lambda: self.identity["value"]
        self.app = app
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        self.app.dependency_overrides.clear()
        asyncio.run(self.engine.dispose())

    def test_dashboard_authorized_returns_200(self):
        response = self.client.get(
            f"/api/v1/teacher/dashboard?school_id={self.school_a}&classroom_id=TURMA_3A"
        )
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["teacher_id"], "teacher-a")
        self.assertEqual(body["student_count"], 1)

    def test_dashboard_unauthorized_classroom_returns_403(self):
        response = self.client.get(
            f"/api/v1/teacher/dashboard?school_id={self.school_b}&classroom_id=TURMA_UNKNOWN"
        )
        self.assertEqual(response.status_code, 403, response.text)
        self.assertIn("detail", response.json())

    def test_list_classrooms_authorized_returns_200(self):
        response = self.client.get(f"/api/v1/teacher/classrooms?school_id={self.school_a}")
        self.assertEqual(response.status_code, 200, response.text)
        classroom_ids = [item["classroom_id"] for item in response.json()]
        self.assertIn("TURMA_3A", classroom_ids)

    def test_list_classrooms_unauthorized_school_returns_empty_list_not_403(self):
        # list_teacher_classrooms never raises ScopeAuthorizationError - a
        # teacher with zero authorization in school_id legitimately gets an
        # empty scope, not a denial. Confirms the router's except branch is
        # simply unreachable here (unlike dashboard/detail/export), while
        # asserting the actual documented behavior stays a real assertion.
        response = self.client.get(f"/api/v1/teacher/classrooms?school_id={self.school_b}")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json(), [])

    def test_classroom_detail_authorized_returns_200(self):
        response = self.client.get(
            f"/api/v1/teacher/classrooms/TURMA_3A?school_id={self.school_a}"
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("action_plan", response.json())

    def test_classroom_detail_unauthorized_classroom_returns_403(self):
        response = self.client.get(
            f"/api/v1/teacher/classrooms/TURMA_UNKNOWN?school_id={self.school_a}"
        )
        self.assertEqual(response.status_code, 403, response.text)

    def test_student_detail_authorized_returns_200(self):
        response = self.client.get(
            f"/api/v1/teacher/students/student-a?school_id={self.school_a}"
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["student_id"], "student-a")

    def test_student_detail_unauthorized_returns_403(self):
        response = self.client.get(
            f"/api/v1/teacher/students/student-a?school_id={self.school_b}"
        )
        self.assertEqual(response.status_code, 403, response.text)

    def test_search_students_authorized_returns_200(self):
        response = self.client.get(
            f"/api/v1/teacher/search?q=student&school_id={self.school_a}"
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()[0]["student_id"], "student-a")

    def test_search_students_unauthorized_school_returns_empty_list_not_403(self):
        # Same as list_teacher_classrooms: search_students_in_scope resolves
        # to an empty authorized-classroom set rather than raising, so the
        # route's 403 branch is unreachable for this case too.
        response = self.client.get(
            f"/api/v1/teacher/search?q=student&school_id={self.school_b}"
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json(), [])

    def test_export_classroom_report_pdf_returns_200(self):
        # The route used to lie: response_model=ReportExportResponse forced a
        # JSON body while content_type claimed "application/pdf" - no real
        # file ever reached the browser. This asserts on ACTUAL PDF bytes.
        response = self.client.get(
            f"/api/v1/teacher/classrooms/TURMA_3A/export?school_id={self.school_a}&format=pdf"
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers["content-type"], "application/pdf")
        self.assertIn("attachment", response.headers["content-disposition"])
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertGreater(len(response.content), 500)

        import pymupdf
        doc = pymupdf.open(stream=response.content, filetype="pdf")
        try:
            full_text = "".join(page.get_text() for page in doc)
        finally:
            doc.close()
        self.assertIn("TURMA_3A", full_text)

    def test_export_classroom_report_xlsx_returns_200(self):
        response = self.client.get(
            f"/api/v1/teacher/classrooms/TURMA_3A/export?school_id={self.school_a}&format=xlsx"
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(
            response.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment", response.headers["content-disposition"])
        self.assertTrue(response.content.startswith(b"PK"))

        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(response.content))
        self.assertGreaterEqual(len(wb.sheetnames), 1)

    def test_export_classroom_report_invalid_format_returns_400(self):
        response = self.client.get(
            f"/api/v1/teacher/classrooms/TURMA_3A/export?school_id={self.school_a}&format=docx"
        )
        self.assertEqual(response.status_code, 400, response.text)
        self.assertIn("Unsupported export format", response.json()["detail"])

    def test_export_classroom_report_unauthorized_classroom_returns_403(self):
        response = self.client.get(
            f"/api/v1/teacher/classrooms/TURMA_UNKNOWN/export?school_id={self.school_a}&format=pdf"
        )
        self.assertEqual(response.status_code, 403, response.text)


if __name__ == "__main__":
    unittest.main()
