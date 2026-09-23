"""HTTP-layer tests for authorial_ingestion.py (src/agente_ia_edu/api/routes/authorial_ingestion.py).

Wave-2 (overnight HTTP-route-coverage campaign, following 703661f): before
this file, list_ingestions, update_classification, reject_ingestion, and
publish_ingestion had ZERO HTTP-level coverage (test_phase26_authorial_
material_ingestion.py exercises the service layer directly, and
test_r0_authorial_ingestion_scope.py unit-tests _require_review_access in
isolation) - every existing test_tenant_isolation_via_api()/
test_student_cannot_upload() HTTP call in test_phase26_* only ever hits the
403 branches of upload/get/approve, never their success paths or the other
four endpoints at all.

Also covers the upload endpoint's own error-mapping branches (unsupported
format -> 422, oversized file -> 413, AuthorialIngestionError -> 422 via
_map_error) which no test had reached either.
"""

from __future__ import annotations

import asyncio
import unittest
import uuid as _uuid
from pathlib import Path

from fastapi.testclient import TestClient

from agente_ia_edu.api.app import create_app
from agente_ia_edu.api.dependencies import get_current_identity, get_session_factory
from agente_ia_edu.db.base import Base
from agente_ia_edu.db.models import CatalogNode, IngestionDocument
from agente_ia_edu.db.models.admin import School, UserSchoolLink
from agente_ia_edu.identity import ExternalIdentityContext
from agente_ia_edu.services.authorial_material_ingestion import AuthorialMaterialIngestionService
from agente_ia_edu.services.material_storage import MaterialStorage, file_sha256
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

_SCHOOL_A = _uuid.uuid5(_uuid.NAMESPACE_DNS, "authih-school-a")
_SCHOOL_B = _uuid.uuid5(_uuid.NAMESPACE_DNS, "authih-school-b")


def _ident(user: str, *, roles: tuple[str, ...] = ("teacher",)) -> ExternalIdentityContext:
    return ExternalIdentityContext(provider="test", external_user_id=user, roles=roles)


def _make_docx(path: Path, *, title: str = "Apostila Teste HTTP") -> None:
    from docx import Document
    doc = Document()
    doc.add_paragraph(f"Episódio 01 – {title}")
    doc.add_paragraph(f"Um parágrafo de introdução sobre o tema estudado ({path.stem}).")
    doc.add_paragraph("Questão 1. Qual é a resposta correta?")
    doc.add_paragraph("a) Alternativa um")
    doc.add_paragraph("b) Alternativa dois")
    doc.save(path)


async def _seed(factory) -> None:
    async with factory() as s:
        d = CatalogNode(code="AIH-DISC", name="Disciplina HTTP", node_type="DISCIPLINE", active=True)
        s.add(d); await s.flush(); d.root_id = d.id
        a = CatalogNode(code="AIH-AREA", name="Area HTTP", node_type="AREA",
                        parent_id=d.id, root_id=d.id, active=True)
        s.add(a); await s.flush()
        c = CatalogNode(code="AIH-CONTENT", name="Apostila Teste HTTP", node_type="CONTENT",
                        parent_id=a.id, root_id=d.id, active=True)
        s.add(c); await s.flush()
        s.add(School(id=_SCHOOL_A, code="AIHSCHA", name="Escola AIH A", status="ACTIVE"))
        s.add(School(id=_SCHOOL_B, code="AIHSCHB", name="Escola AIH B", status="ACTIVE"))
        await s.flush()
        s.add(UserSchoolLink(external_user_id="prof_a", school_id=_SCHOOL_A, role="TEACHER",
                             scope_type="SCHOOL", scope_external_id=str(_SCHOOL_A), active=True))
        s.add(UserSchoolLink(external_user_id="prof_b", school_id=_SCHOOL_B, role="TEACHER",
                             scope_type="SCHOOL", scope_external_id=str(_SCHOOL_B), active=True))
        await s.commit()


class AuthorialIngestionHTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.loop = asyncio.new_event_loop()
        cls.engine = create_async_engine("sqlite+aiosqlite:///:memory:", poolclass=StaticPool)
        cls.factory = async_sessionmaker(cls.engine, class_=AsyncSession, expire_on_commit=False)

        async def _prep():
            async with cls.engine.begin() as conn:
                await conn.run_sync(Base.metadata.create_all)
            await _seed(cls.factory)

        cls.loop.run_until_complete(_prep())

        cls.tmp = Path("/tmp/authorial_ingestion_http_fixtures")
        cls.tmp.mkdir(exist_ok=True)
        cls.storage_root = Path("/tmp/authorial_ingestion_http_storage")

        cls.app = create_app()
        cls.app.dependency_overrides[get_session_factory] = lambda: cls.factory
        cls.app.dependency_overrides[get_current_identity] = lambda: _ident("prof_a")
        cls.client = TestClient(cls.app)

    @classmethod
    def tearDownClass(cls):
        cls.app.dependency_overrides.clear()
        cls.loop.run_until_complete(cls.engine.dispose())
        cls.loop.close()

    def setUp(self):
        self._as("prof_a")

    def _as(self, user: str, **kwargs) -> None:
        self.app.dependency_overrides[get_current_identity] = lambda: _ident(user, **kwargs)

    def _svc(self, session) -> AuthorialMaterialIngestionService:
        return AuthorialMaterialIngestionService(session, storage=MaterialStorage(root=self.storage_root))

    def _upload(self, path: Path, *, filename: str | None = None):
        with open(path, "rb") as f:
            return self.client.post(
                "/api/v1/catalog/ingestion/upload",
                files={"file": (
                    filename or path.name, f,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )},
            )

    def _upload_mapped(self, name: str) -> dict:
        """Upload a doc whose title matches the seeded CONTENT node exactly
        (classification_state == MAPPED, review_status == PENDING_REVIEW)."""
        path = self.tmp / name
        _make_docx(path, title="Apostila Teste HTTP")
        r = self._upload(path)
        self.assertEqual(r.status_code, 201, r.text)
        return r.json()

    def _upload_gap(self, name: str) -> dict:
        """Upload a doc with an unrelated title (classification_state ==
        TAXONOMY_GAP, review_status == NEEDS_REVIEW)."""
        path = self.tmp / name
        _make_docx(path, title="Zzqxw Totally Unrelated Whatever 999")
        r = self._upload(path)
        self.assertEqual(r.status_code, 201, r.text)
        return r.json()

    # -- upload: format/size/duplicate error-mapping branches --------------

    def test_upload_unsupported_format_returns_422(self):
        path = self.tmp / "notes.png"
        path.write_bytes(b"\x89PNG\r\n\x1a\nnot a real png but bytes")
        with open(path, "rb") as f:
            r = self.client.post(
                "/api/v1/catalog/ingestion/upload",
                files={"file": ("notes.png", f, "image/png")},
            )
        self.assertEqual(r.status_code, 422, r.text)
        self.assertIn("unsupported file format", r.json()["detail"])

    def test_upload_oversized_file_returns_413(self):
        path = self.tmp / "huge.txt"
        # 26MB of content > the router's 25MB ceiling, forcing the chunked
        # read loop's size check (never an unbounded upload).
        with open(path, "wb") as f:
            chunk = b"x" * (1024 * 1024)
            for _ in range(26):
                f.write(chunk)
        with open(path, "rb") as f:
            r = self.client.post(
                "/api/v1/catalog/ingestion/upload",
                files={"file": ("huge.txt", f, "text/plain")},
            )
        self.assertEqual(r.status_code, 413, r.text)
        self.assertIn("25MB", r.json()["detail"])

    def test_upload_hash_collision_outside_authorial_workflow_returns_422(self):
        # A document already ingested through the OFFICIAL (non-authorial)
        # pipeline, with no IngestionMaterialReview attached to it -
        # ingest_file() must fail closed (never silently adopt it), which
        # exercises both _map_error's AuthorialIngestionError branch and the
        # upload route's own except/_map_error call.
        path = self.tmp / "collision.docx"
        _make_docx(path, title="Collision Source")
        digest = file_sha256(path)

        async def seed_official_doc():
            async with self.factory() as s:
                s.add(IngestionDocument(
                    filename="collision.docx", document_type="DOCX", document_hash=digest,
                    storage_uri="file:///dev/null", file_size_bytes=path.stat().st_size,
                    status="processed",
                ))
                await s.commit()
        self.loop.run_until_complete(seed_official_doc())

        r = self._upload(path)
        self.assertEqual(r.status_code, 422, r.text)
        detail = r.json()["detail"]
        self.assertEqual(detail["code"], "EXTRACTION_FAILED")
        self.assertIn("already ingested outside", detail["message"])
        self.assertIn("ingestion_document_id", detail)

    # -- get: not-found mapping ---------------------------------------------

    def test_get_ingestion_not_found_returns_404(self):
        r = self.client.get(f"/api/v1/catalog/ingestion/{_uuid.uuid4()}")
        self.assertEqual(r.status_code, 404, r.text)

    def test_get_ingestion_success_returns_full_detail(self):
        review = self._upload_mapped("get_detail.docx")
        review_id = review["review"]["id"]
        r = self.client.get(f"/api/v1/catalog/ingestion/{review_id}")
        self.assertEqual(r.status_code, 200, r.text)
        body = r.json()
        self.assertEqual(body["review"]["id"], review_id)
        self.assertIn("document", body)
        self.assertIn("sections", body)
        self.assertIn("exercises", body)
        self.assertGreaterEqual(len(body["exercises"]), 1)

    # -- list: scoping (BUG: self-asserted role with no real school link ----
    # currently leaks every school's reviews - see fix in the route file) ---

    def test_list_ingestions_scoped_to_callers_school(self):
        review_a = self._upload_mapped("list_scope_a.docx")
        self._as("prof_b")
        review_b = self._upload_mapped("list_scope_b.docx")
        self._as("prof_a")

        listing = self.client.get("/api/v1/catalog/ingestion")
        self.assertEqual(listing.status_code, 200, listing.text)
        ids = {item["id"] for item in listing.json()}
        self.assertIn(review_a["review"]["id"], ids)
        self.assertNotIn(review_b["review"]["id"], ids)

    def test_list_ingestions_review_status_filter(self):
        mapped = self._upload_mapped("list_filter_mapped.docx")
        gap = self._upload_gap("list_filter_gap.docx")

        needs_review = self.client.get(
            "/api/v1/catalog/ingestion", params={"review_status": "NEEDS_REVIEW"}
        )
        self.assertEqual(needs_review.status_code, 200, needs_review.text)
        ids = {item["id"] for item in needs_review.json()}
        self.assertIn(gap["review"]["id"], ids)
        self.assertNotIn(mapped["review"]["id"], ids)

    def test_list_ingestions_self_asserted_role_without_school_link_is_scoped_empty(self):
        # prof_a (real link, school A) and prof_b (real link, school B) each
        # upload a document. A caller who merely CLAIMS the TEACHER role via
        # an Authorization header, with no real UserSchoolLink at all (the
        # exact fallback AuthorizationService.resolve_context documents -
        # role=self-asserted, school_id=None), must never see either
        # school's reviews. This is the identical bug class already fixed
        # for get/approve/reject/publish/update_classification via
        # _require_review_access (see test_r0_authorial_ingestion_scope.py)
        # - list_ingestions had no equivalent gate at all.
        self._upload_mapped("ghost_scope_a.docx")
        self._as("prof_b")
        self._upload_mapped("ghost_scope_b.docx")

        self._as("ghost_teacher", roles=("teacher",))
        listing = self.client.get("/api/v1/catalog/ingestion")
        self.assertEqual(listing.status_code, 200, listing.text)
        self.assertEqual(listing.json(), [])
        self._as("prof_a")

    # -- update_classification: success + PUBLICATION_BLOCKED ---------------

    def test_update_classification_resolves_taxonomy_gap(self):
        review = self._upload_gap("classify_gap.docx")
        review_id = review["review"]["id"]
        self.assertEqual(review["review"]["classification_state"], "TAXONOMY_GAP")

        r = self.client.patch(
            f"/api/v1/catalog/ingestion/{review_id}/classification",
            json={"content_code": "AIH-CONTENT"},
        )
        self.assertEqual(r.status_code, 200, r.text)
        body = r.json()
        self.assertEqual(body["classification_state"], "MAPPED")
        self.assertEqual(body["classification_source"], "MANUAL")
        self.assertEqual(body["review_status"], "PENDING_REVIEW")

    def test_update_classification_after_publish_is_blocked(self):
        review_id = self._publish_flow("classify_after_publish.docx")
        r = self.client.patch(
            f"/api/v1/catalog/ingestion/{review_id}/classification",
            json={"notes": "tentativa de editar depois de publicado"},
        )
        self.assertEqual(r.status_code, 422, r.text)
        self.assertEqual(r.json()["detail"]["code"], "PUBLICATION_BLOCKED")

    # -- reject ---------------------------------------------------------------

    def test_reject_ingestion_success(self):
        review = self._upload_mapped("reject_me.docx")
        review_id = review["review"]["id"]
        r = self.client.post(
            f"/api/v1/catalog/ingestion/{review_id}/reject",
            json={"reason": "conteúdo desatualizado"},
        )
        self.assertEqual(r.status_code, 200, r.text)
        body = r.json()
        self.assertEqual(body["review_status"], "REJECTED")
        self.assertEqual(body["notes"], "conteúdo desatualizado")

    def test_reject_ingestion_not_found_returns_404(self):
        r = self.client.post(
            f"/api/v1/catalog/ingestion/{_uuid.uuid4()}/reject", json={"reason": "x"}
        )
        self.assertEqual(r.status_code, 404, r.text)

    # -- approve success (HTTP 200 body, not just the 403 cross-tenant path) -

    def test_approve_ingestion_success_returns_body(self):
        review = self._upload_mapped("approve_me.docx")
        review_id = review["review"]["id"]
        r = self.client.post(f"/api/v1/catalog/ingestion/{review_id}/approve")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["review_status"], "APPROVED")

    # -- publish: success + PUBLICATION_BLOCKED before approval -------------

    def _publish_flow(self, name: str) -> str:
        review = self._upload_mapped(name)
        review_id = review["review"]["id"]
        approve = self.client.post(f"/api/v1/catalog/ingestion/{review_id}/approve")
        self.assertEqual(approve.status_code, 200, approve.text)
        publish = self.client.post(f"/api/v1/catalog/ingestion/{review_id}/publish")
        self.assertEqual(publish.status_code, 200, publish.text)
        body = publish.json()
        self.assertEqual(body["review_status"], "PUBLISHED")
        self.assertIsNotNone(body["theory_material_id"])
        return review_id

    def test_publish_ingestion_success_creates_theory_material(self):
        self._publish_flow("publish_me.docx")

    def test_publish_without_approval_is_blocked(self):
        review = self._upload_mapped("publish_no_approval.docx")
        review_id = review["review"]["id"]
        r = self.client.post(f"/api/v1/catalog/ingestion/{review_id}/publish")
        self.assertEqual(r.status_code, 422, r.text)
        self.assertEqual(r.json()["detail"]["code"], "PUBLICATION_BLOCKED")


if __name__ == "__main__":
    unittest.main()
