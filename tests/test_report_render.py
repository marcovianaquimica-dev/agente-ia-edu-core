"""Unit tests for services.report_render - the binary PDF/XLSX renderers that
turn a ReportExportService payload dict into REAL bytes (PHASE 12B.2/12C.2
export-delivery fix: the routes used to serialize this same payload as JSON
while lying about content_type). These tests open the produced bytes with
PyMuPDF / openpyxl and assert on real content, not just "it didn't crash".
"""

from __future__ import annotations

import io
import unittest

from agente_ia_edu.services.report_export import ReportExportService


def _classroom_payload(fmt: str) -> dict:
    classroom_data = {
        "classroom_id": "TURMA_3A",
        "summary": {
            "student_count": 30,
            "active_students_count": 28,
            "overall_class_average": 61.5,
        },
        "mastery_distribution": {
            "struggling_percentage": 20.0,
            "developing_percentage": 45.0,
            "mastered_percentage": 35.0,
        },
        "strengths": [
            {"content_node_id": "c1", "content_name": "Estequiometria",
             "class_average_mastery": 82.0, "reason": "Alta média"},
        ],
        "improvement_areas": [
            {"content_node_id": "c2", "content_name": "Termoquímica",
             "class_average_mastery": 41.0, "students_struggling_count": 9,
             "reason": "Baixa média"},
        ],
        "recent_contents_taught": [
            {"content_node_id": "c2", "content_name": "Termoquímica",
             "last_lesson_date": "2026-09-10T00:00:00+00:00",
             "teacher_or_author": "teacher-a", "class_average_mastery": 41.0,
             "struggling_students_count": 9,
             "recommended_action": "Revisão conceitual urgente."},
        ],
        "action_plan": [
            {"priority": "HIGH", "content_name": "Termoquímica",
             "class_average_mastery": 41.0, "impacted_students_count": 9,
             "evidence": "9 alunos abaixo de 50%",
             "recommended_action": "Revisão conceitual urgente."},
        ],
        "students": [
            {"student_id": "student-a", "name": "Aluno A", "classroom_id": "TURMA_3A",
             "average_mastery": 55.0, "status_label": "Em desenvolvimento"},
        ],
    }
    return ReportExportService.export_classroom_report(classroom_data, export_format=fmt)


def _student_payload(fmt: str) -> dict:
    student_data = {
        "student_id": "student-a",
        "accuracy_percentage": 72.5,
        "total_questions_answered": 40,
        "content_masteries": [
            {"content_node_id": "c1", "content_name": "Estequiometria",
             "mastery_score": 82.0, "current_level": "CONSOLIDATED",
             "questions_answered": 10, "questions_correct": 8},
        ],
        "priority_contents": [
            {"content_node_id": "c2", "content_name": "Termoquímica",
             "mastery_score": 41.0, "current_level": "STRUGGLING",
             "questions_answered": 6, "questions_correct": 2},
        ],
        "current_recommendations": [
            {"content_node_id": "c2", "content_name": "Termoquímica",
             "reason": "Domínio abaixo de 50%"},
        ],
    }
    return ReportExportService.export_student_report(student_data, export_format=fmt)


class ReportRenderPDF(unittest.TestCase):
    def test_classroom_pdf_is_real_pdf_with_expected_text(self):
        from agente_ia_edu.services.report_render import render_pdf

        payload = _classroom_payload("pdf")
        data = render_pdf(payload)

        self.assertTrue(data.startswith(b"%PDF-"))
        self.assertGreater(len(data), 500)

        import pymupdf
        doc = pymupdf.open(stream=data, filetype="pdf")
        try:
            self.assertGreaterEqual(len(doc), 1)
            full_text = "".join(page.get_text() for page in doc)
        finally:
            doc.close()

        self.assertIn("Estequiometria", full_text)
        self.assertIn("Termoquímica", full_text)
        self.assertIn("TURMA_3A", full_text)

    def test_student_pdf_is_real_pdf_with_expected_text(self):
        from agente_ia_edu.services.report_render import render_pdf

        payload = _student_payload("pdf")
        data = render_pdf(payload)

        self.assertTrue(data.startswith(b"%PDF-"))

        import pymupdf
        doc = pymupdf.open(stream=data, filetype="pdf")
        try:
            full_text = "".join(page.get_text() for page in doc)
        finally:
            doc.close()

        self.assertIn("student-a", full_text)
        self.assertIn("Estequiometria", full_text)
        self.assertIn("Termoquímica", full_text)


class ReportRenderXLSX(unittest.TestCase):
    def test_classroom_xlsx_is_real_workbook_with_expected_cells(self):
        from agente_ia_edu.services.report_render import render_xlsx

        payload = _classroom_payload("xlsx")
        data = render_xlsx(payload)

        self.assertTrue(data.startswith(b"PK"))

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data))

        sheet_names = wb.sheetnames
        self.assertGreater(len(sheet_names), 1)

        # A "Pontos Fortes" sheet must exist with the strength row rendered.
        strengths_sheet = next(s for s in sheet_names if "Fortes" in s)
        ws = wb[strengths_sheet]
        values = [cell.value for row in ws.iter_rows() for cell in row if cell.value]
        self.assertTrue(any("Estequiometria" in str(v) for v in values))

        roster_sheet = next(s for s in sheet_names if "Alunos" in s)
        ws2 = wb[roster_sheet]
        values2 = [cell.value for row in ws2.iter_rows() for cell in row if cell.value]
        self.assertTrue(any("student-a" in str(v) for v in values2))

    def test_student_xlsx_is_real_workbook_with_expected_cells(self):
        from agente_ia_edu.services.report_render import render_xlsx

        payload = _student_payload("xlsx")
        data = render_xlsx(payload)

        self.assertTrue(data.startswith(b"PK"))

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data))
        all_values = [
            cell.value
            for sheet in wb.sheetnames
            for row in wb[sheet].iter_rows()
            for cell in row
            if cell.value
        ]
        self.assertTrue(any("Termoquímica" in str(v) for v in all_values))
        self.assertTrue(any(v == 72.5 or v == "72.5" for v in all_values))


if __name__ == "__main__":
    unittest.main()
