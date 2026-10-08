"""Gera a amostra CEGA de adjudicacao humana, V1 x V2.

PROTOTIPO EXPERIMENTAL. Nao abre banco, nao chama IA, nao baixa nada.

Cegueira: a planilha NAO diz de qual versao veio cada reconstrucao. O mapa
``id -> origem`` e gravado num arquivo SEPARADO, que o adjudicador nao abre.
Sem isso a adjudicacao mede a expectativa do adjudicador, nao a fidelidade da
extracao.

Amostragem deterministica: ``random.Random(SEMENTE)`` com semente fixa e
declarada. Rodar duas vezes produz a mesma amostra.

NENHUM JULGAMENTO E PREENCHIDO. As colunas A-G saem vazias.

Uso:
    python tests/manual/enem_v2_amostra_cega.py --saida /caminho/amostra.xlsx
"""

from __future__ import annotations

import argparse
import json
import random
import re
import sys
from pathlib import Path

_AQUI = Path(__file__).resolve().parent
for _cand in (Path.cwd(), _AQUI.parents[1]):
    if (_cand / "src" / "agente_ia_edu").is_dir():
        sys.path.insert(0, str(_cand / "src"))
        RAIZ = _cand
        break
else:  # pragma: no cover
    RAIZ = _AQUI.parents[1]
    sys.path.insert(0, str(RAIZ / "src"))

from agente_ia_edu.services.enem_extraction_v2 import extrair_caderno  # noqa: E402
from agente_ia_edu.services.ingestion_parser import PdfParser  # noqa: E402

SEMENTE = 20261003
POR_ESTRATO = 5
MANIFESTO = RAIZ / "tests/manual/phase10_enem_manifest_2016_2025.json"

ESTRATOS = (
    ("E1_AMBAS", "reconstruida pelas duas versoes"),
    ("E2_SO_V2", "recuperada apenas pela versao experimental"),
    ("E3_SO_V1", "reconstruida apenas pela versao baseline (regressao)"),
    ("E4_COM_ATIVO", "item com elemento visual identificado"),
    ("E5_SEM_ATIVO", "item sem elemento visual identificado"),
    ("E6_QUEBRA_PAGINA", "item que atravessa quebra de pagina"),
    ("E7_ERA_REVISAO", "item que o baseline mandava para revisao"),
    ("E8_FORMA_RARA", "forma de marcador de alternativa menos frequente"),
)

CRITERIOS = (
    ("A_numero", "O numero do item esta correto?"),
    ("B_enunciado", "O enunciado esta completo, sem faltar trecho?"),
    ("C_alternativas", "As cinco alternativas estao completas?"),
    ("D_ordem", "A ordem A-B-C-D-E esta correta?"),
    ("E_sem_lixo", "Esta livre de texto estranho agregado (rodape, marca d'agua)?"),
    ("F_visual", "O elemento visual necessario foi identificado?"),
    ("G_gabarito", "O gabarito esta corretamente associado?"),
)


# O Excel recusa caracteres de controle. O texto degradado de 2021 os contem.
# Substituir por um marcador visivel preserva a INFORMACAO de que havia lixo
# ali - apagar silenciosamente esconderia do adjudicador justamente o defeito
# que ele precisa ver.
_CONTROLE = re.compile(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]")


def para_planilha(texto: str, limite: int = 2000) -> str:
    return _CONTROLE.sub("\u00a4", texto or "")[:limite]


def coletar(ano: int, dia: int, entrada: dict) -> list[dict]:
    prova, gabarito = RAIZ / entrada["proof_pdf"], RAIZ / entrada["answer_key_pdf"]

    analisado = PdfParser.parse_file(prova)
    chave_v1 = PdfParser.parse_answer_key(gabarito)
    PdfParser.apply_answer_key(analisado, chave_v1)
    v1 = {}
    for q in analisado.questions:
        v1.setdefault(q.question_number, {
            "numero": q.question_number, "enunciado": q.statement_text,
            "alternativas": q.alternatives_text or "", "gabarito": q.correct_answer or "",
            "pagina": q.page_start, "revisao": q.requires_review,
            "ativos": sum(1 for a in analisado.assets
                          if a.question_number == q.question_number),
            "atravessa": "", "forma": "", "origem": "V1",
        })

    r = extrair_caderno(prova, gabarito, ano=ano, dia=dia, caderno=entrada["booklet"])
    v2 = {}
    for q in r.questoes:
        if q.numero in v2:
            continue
        v2[q.numero] = {
            "numero": q.numero, "enunciado": q.enunciado,
            "alternativas": q.alternativas_texto or "", "gabarito": q.gabarito or "",
            "pagina": q.pagina_inicio, "revisao": bool(q.problemas),
            "ativos": len(q.ativos), "atravessa": "sim" if q.atravessa_pagina else "",
            "forma": q.forma_alternativas, "origem": "V2",
            "problemas": ";".join(q.problemas),
        }

    def reconstruiu(d: dict) -> bool:
        return bool(d.get("alternativas")) and len(d["alternativas"].splitlines()) == 5

    linhas = []
    for numero in sorted(set(v1) | set(v2)):
        a, b = v1.get(numero), v2.get(numero)
        ok1, ok2 = bool(a and reconstruiu(a)), bool(b and reconstruiu(b))
        estratos = []
        if ok1 and ok2:
            estratos.append("E1_AMBAS")
        if ok2 and not ok1:
            estratos.append("E2_SO_V2")
        if ok1 and not ok2:
            estratos.append("E3_SO_V1")
        alvo = b if ok2 else a
        if alvo is None:
            continue
        if alvo["ativos"]:
            estratos.append("E4_COM_ATIVO")
        else:
            estratos.append("E5_SEM_ATIVO")
        if b and b["atravessa"]:
            estratos.append("E6_QUEBRA_PAGINA")
        if a and a["revisao"]:
            estratos.append("E7_ERA_REVISAO")
        if b and b["forma"] in ("LETRA_SOZINHA", "2+ESPACOS", "TAB"):
            estratos.append("E8_FORMA_RARA")
        for candidato, veio in ((a, "V1"), (b, "V2")):
            if candidato is None:
                continue
            linhas.append({**candidato, "ano": ano, "dia": dia,
                           "caderno": entrada["booklet"], "origem": veio,
                           "estratos": estratos})
    return linhas


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--saida", required=True)
    p.add_argument("--por-estrato", type=int, default=POR_ESTRATO,
                   help="linhas por estrato; aumentar da estimativa mais fina")
    args = p.parse_args()
    por_estrato = args.por_estrato
    destino = Path(args.saida)

    manifesto = {(e["exam_year"], e["exam_day"]): e
                 for e in json.loads(MANIFESTO.read_text())["booklets"]}

    todas = []
    for ano in range(2020, 2026):
        for dia in (1, 2):
            print(f"... {ano} D{dia}", flush=True)
            todas.extend(coletar(ano, dia, manifesto[(ano, dia)]))

    aleatorio = random.Random(SEMENTE)
    escolhidas: dict[tuple, dict] = {}
    cobertura: dict[str, int] = {}
    for nome, _descricao in ESTRATOS:
        candidatas = [linha for linha in todas if nome in linha["estratos"]]
        # espalha por ano e dia antes de sortear
        por_caderno: dict[tuple, list] = {}
        for linha in candidatas:
            por_caderno.setdefault((linha["ano"], linha["dia"]), []).append(linha)
        sorteadas = []
        for caderno in sorted(por_caderno):
            aleatorio.shuffle(por_caderno[caderno])
            sorteadas.extend(por_caderno[caderno][:1])
        aleatorio.shuffle(sorteadas)
        for linha in sorteadas[:por_estrato]:
            escolhidas[(linha["ano"], linha["dia"], linha["numero"], linha["origem"])] = linha
        cobertura[nome] = min(len(sorteadas), por_estrato)

    amostra = list(escolhidas.values())
    aleatorio.shuffle(amostra)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:  # pragma: no cover
        print("openpyxl ausente")
        return

    livro = Workbook()

    folha = livro.active
    folha.title = "INSTRUCOES"
    texto = [
        ("ADJUDICACAO CEGA - RECONSTRUCAO DE ITEM DO ENEM", True),
        ("", False),
        ("ARQUIVO RESTRITO. Contem texto de prova oficial do INEP.", True),
        ("", False),
        ("O que fazer:", True),
        ("1. Abra o PDF do caderno indicado em cada linha (var/inep-pilot/).", False),
        ("2. Va a pagina indicada e encontre o item pelo numero.", False),
        ("3. Compare o que esta impresso com a reconstrucao da planilha.", False),
        ("4. Preencha as colunas A a G com SIM, NAO ou NA.", False),
        ("   NA = nao se aplica (ex.: item sem elemento visual na coluna F).", False),
        ("5. Use OBSERVACAO para o que nao couber em SIM/NAO.", False),
        ("", False),
        ("Importante:", True),
        ("- A planilha NAO diz de qual versao do extrator veio cada linha.", False),
        ("  Isso e proposital. Nao tente adivinhar.", False),
        ("- O mesmo item pode aparecer duas vezes, em linhas separadas e", False),
        ("  distantes. Julgue cada linha por si, sem procurar a outra.", False),
        ("- Julgue o que esta escrito, nao o que deveria estar.", False),
        ("- Em duvida, escreva a duvida em OBSERVACAO em vez de chutar.", False),
        ("", False),
        ("Os sete criterios:", True),
    ]
    for chave, pergunta in CRITERIOS:
        texto.append((f"   {chave}: {pergunta}", False))
    texto += [
        ("", False),
        (f"Amostra: {len(amostra)} linhas, semente {SEMENTE}, deterministica.", False),
        ("Estratos cobertos:", True),
    ]
    for nome, descricao in ESTRATOS:
        texto.append((f"   {nome} ({cobertura.get(nome, 0)} linhas): {descricao}", False))

    for indice, (linha_texto, negrito) in enumerate(texto, start=1):
        celula = folha.cell(row=indice, column=1, value=linha_texto)
        celula.font = Font(bold=negrito)
    folha.column_dimensions["A"].width = 95

    folha = livro.create_sheet("AMOSTRA")
    cabecalho = (["id", "ano", "dia", "caderno", "pagina", "numero_lido",
                  "enunciado_reconstruido", "alternativas_reconstruidas",
                  "gabarito_associado", "elemento_visual_detectado"]
                 + [c for c, _ in CRITERIOS] + ["OBSERVACAO"])
    folha.append(cabecalho)
    fundo = PatternFill("solid", fgColor="FFF2CC")
    for coluna in range(1, len(cabecalho) + 1):
        celula = folha.cell(row=1, column=coluna)
        celula.font = Font(bold=True)
        celula.alignment = Alignment(wrap_text=True, vertical="top")
        if cabecalho[coluna - 1] in {c for c, _ in CRITERIOS} | {"OBSERVACAO"}:
            celula.fill = fundo

    mapa = []
    for indice, linha in enumerate(amostra, start=1):
        identificador = f"A{indice:03d}"
        folha.append([
            identificador, linha["ano"], linha["dia"], linha["caderno"],
            linha["pagina"], linha["numero"],
            para_planilha(linha["enunciado"]), para_planilha(linha["alternativas"]),
            linha["gabarito"] or "(nenhum)",
            "sim" if linha["ativos"] else "nao",
        ] + [""] * (len(CRITERIOS) + 1))
        mapa.append({"id": identificador, "origem": linha["origem"],
                     "ano": linha["ano"], "dia": linha["dia"],
                     "numero": linha["numero"], "estratos": linha["estratos"],
                     "problemas": linha.get("problemas", "")})

    for coluna, largura in (("A", 7), ("B", 6), ("C", 5), ("D", 9), ("E", 8),
                            ("F", 12), ("G", 62), ("H", 62), ("I", 10), ("J", 12)):
        folha.column_dimensions[coluna].width = largura
    for coluna in "KLMNOPQR":
        folha.column_dimensions[coluna].width = 13
    for linha_idx in range(2, len(amostra) + 2):
        for coluna in range(1, len(cabecalho) + 1):
            folha.cell(row=linha_idx, column=coluna).alignment = Alignment(
                wrap_text=True, vertical="top")

    livro.save(destino)
    chave = destino.with_name(destino.stem + "_CHAVE_NAO_ABRIR.json")
    chave.write_text(json.dumps(
        {"semente": SEMENTE, "n": len(amostra), "mapa": mapa},
        ensure_ascii=False, indent=1))

    print(f"\nplanilha : {destino}  ({len(amostra)} linhas)")
    print(f"chave    : {chave}   <- NAO abrir antes de preencher")
    print("\ncobertura por estrato:")
    for nome, descricao in ESTRATOS:
        print(f"   {nome:<18} {cobertura.get(nome, 0):>2}  {descricao}")
    v1 = sum(1 for m in mapa if m["origem"] == "V1")
    print(f"\norigem (so na chave): V1={v1}  V2={len(mapa) - v1}")


if __name__ == "__main__":
    main()
