import io
import unittest
import uuid

import openpyxl

from agente_ia_edu.services.essay_dashboard_export import build_essay_dashboard_xlsx
from agente_ia_edu.services.essay_teacher_dashboard import EssayDashboardResponse, StudentSubmissionRow


def _dashboard() -> EssayDashboardResponse:
    return EssayDashboardResponse(
        essay_prompt_id=uuid.uuid4(), essay_prompt_title="Tema",
        total_students=2, submitted_count=1, submitted_percentage=50.0,
        average_total_score=640.0,
        average_per_competency={"C1": 128.0, "C2": 128.0, "C3": 128.0, "C4": 128.0, "C5": 128.0},
        students=[
            StudentSubmissionRow(
                student_id=uuid.uuid4(), student_name="Aluno A", class_id=uuid.uuid4(),
                submitted=True, total_score=640,
                per_competency={"C1": 128, "C2": 128, "C3": 128, "C4": 128, "C5": 128},
            ),
            StudentSubmissionRow(
                student_id=uuid.uuid4(), student_name="Aluno B", class_id=uuid.uuid4(),
                submitted=False,
            ),
        ],
        action_plan=[],
    )


class EssayDashboardExportTests(unittest.TestCase):
    def test_unknown_report_type_raises(self):
        with self.assertRaises(ValueError):
            build_essay_dashboard_xlsx(_dashboard(), report_type="not_a_real_type")

    def test_submission_list_has_name_and_yes_no_columns(self):
        data = build_essay_dashboard_xlsx(_dashboard(), report_type="submission_list")
        wb = openpyxl.load_workbook(io.BytesIO(data))
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Aluno", "Entregou"))
        self.assertIn(("Aluno A", "Sim"), rows)
        self.assertIn(("Aluno B", "Não"), rows)

    def test_grades_total_includes_only_total_score(self):
        data = build_essay_dashboard_xlsx(_dashboard(), report_type="grades_total")
        wb = openpyxl.load_workbook(io.BytesIO(data))
        rows = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Aluno", "Entregou", "Nota total"))
        self.assertIn(("Aluno A", "Sim", 640), rows)
        # Did not submit - score cell must be blank, not a fabricated 0.
        self.assertIn(("Aluno B", "Não", None), rows)

    def test_grades_per_competency_includes_all_five_columns(self):
        data = build_essay_dashboard_xlsx(_dashboard(), report_type="grades_per_competency")
        wb = openpyxl.load_workbook(io.BytesIO(data))
        rows = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Aluno", "Entregou", "C1", "C2", "C3", "C4", "C5", "Nota total"))
        self.assertIn(("Aluno A", "Sim", 128, 128, 128, 128, 128, 640), rows)


if __name__ == "__main__":
    unittest.main()
